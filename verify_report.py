import docx

def verify():
    path = "d:/fsd/Smart_Urban_Green_Management_System_Final_50_Pages_Report.docx"
    doc = docx.Document(path)
    
    page_breaks = 0
    for p in doc.paragraphs:
        for r in p.runs:
            for elem in r._r:
                if elem.tag.endswith('br') and elem.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') == 'page':
                    page_breaks += 1
                    
    total_pages = page_breaks + 1
    total_words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                total_words += len(cell.text.split())

    print(f"File verified: {path}")
    print(f"Total Page Breaks: {page_breaks}")
    print(f"Total Pages: {total_pages}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print(f"Total Word Count: {total_words}")

if __name__ == "__main__":
    verify()
